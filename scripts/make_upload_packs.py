"""Demo file packs for the upload flow: a credit agreement plus management accounts (.xlsx) per company.

Writes demo_files/{nortech,ashby_rowe}/. Nortech: the real March 2024 credit agreement and SEC XBRL quarters up to
30 Sep 2024 (quarterly columns, $k). Ashby & Rowe: the sample agreement and three years of monthly accounts (£k).
"""
import shutil
from datetime import date
from pathlib import Path

from openpyxl import Workbook

from api.cases import DATA, load_case
from scripts.make_sample import LINES

OUT = Path(__file__).resolve().parent.parent / "demo_files"
BALANCES = ("receivables", "inventory", "payables", "cash", "debt")

NORTECH_LINES = [
    ("Net sales", lambda q: q.revenue), ("Cost of goods sold", lambda q: q.cogs),
    ("Selling, general and administrative expense", lambda q: q.revenue - q.cogs - q.ebit - q.da - q.addbacks),
    ("Depreciation and amortization", lambda q: q.da), ("Interest expense", lambda q: q.interest_expense),
    ("Purchases of property and equipment", lambda q: q.capex),
    ("Cash and cash equivalents", lambda q: q.cash), ("Line of credit - Bank of America", lambda q: q.debt),
    ("Accounts receivable, less allowances", lambda q: q.receivables), ("Inventories, net", lambda q: q.inventory),
    ("Accounts payable", lambda q: q.payables),
]


def nortech() -> None:
    d = OUT / "nortech"
    d.mkdir(parents=True, exist_ok=True)
    shutil.copy(DATA / "nortech" / "docs" / "2024-03_credit_agreement.htm", d / "credit_agreement_2024-03.htm")
    qs, _, _ = load_case("nortech")  # cut at as_of, 30 Sep 2024
    wb = Workbook()
    ws = wb.active
    ws.title = "Quarterly accounts"
    ws.append(["Account ($000)"] + [q.period_end.strftime("%b-%y") for q in qs])
    for name, f in NORTECH_LINES:
        ws.append([name] + [round(f(q) * 1000, 1) for q in qs])
    wb.save(d / "nortech_quarterly_accounts.xlsx")


def ashby_rowe() -> None:
    d = OUT / "ashby_rowe"
    d.mkdir(parents=True, exist_ok=True)
    shutil.copy(DATA / "sample_uk" / "agreement.txt", d / "facility_agreement.txt")
    qs, _, _ = load_case("sample_uk")
    wb = Workbook()
    ws = wb.active
    ws.title = "Mgmt accounts"
    ws.append(["Account"] + [date(q.period_end.year, q.period_end.month - 2 + k, 1).strftime("%b-%y")
                             for q in qs for k in range(3)])
    for name, cat, share in LINES + [("Term loan repayments", "principal", 1.0)]:
        row = [name]
        for q in qs:
            for _ in range(3):
                if cat in BALANCES:
                    v = getattr(q, cat) * share
                elif cat == "opex":
                    v = (q.revenue - q.cogs - q.ebitda) * share / 3
                elif cat == "interest":
                    v = q.interest_expense / 3
                elif cat == "principal":
                    v = q.principal_paid / 3
                elif cat == "addbacks":
                    v = 0.0
                else:
                    v = getattr(q, cat) * share / 3
                row.append(round(v * 1000, 1))  # £k
        ws.append(row)
    wb.save(d / "ashby_rowe_management_accounts.xlsx")


def ashby_rowe_budget() -> None:
    """FY27 budget (Oct-26 to Sep-27): last year's quarters +3%, plus a planned £2m machinery spend in Q1 2027."""
    qs, _, _ = load_case("sample_uk")
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget FY27"
    last_year = qs[-4:]
    months = [date(q.period_end.year + 1, q.period_end.month - 2 + k, 1) for q in last_year for k in range(3)]
    ws.append(["Account"] + [m.strftime("%b-%y") for m in months])
    flows = [(n, c, s) for n, c, s in LINES if c not in BALANCES and c != "addbacks"]
    for name, cat, share in flows + [("Term loan repayments", "principal", 1.0)]:
        row = [name]
        for q in last_year:
            for k in range(3):
                if cat == "opex":
                    v = (q.revenue - q.cogs - q.ebitda) * share / 3 * 1.03
                elif cat == "interest":
                    v = q.interest_expense / 3
                elif cat == "principal":
                    v = q.principal_paid / 3
                elif cat == "capex":
                    v = q.capex / 3 + (2.0 / 3 if q.period_end.month == 3 else 0.0)  # new press line, Jan-Mar 27
                elif cat == "da":
                    v = q.da / 3
                else:
                    v = getattr(q, cat) * share / 3 * 1.03
                row.append(round(v * 1000, 1))  # £k
        ws.append(row)
    wb.save(OUT / "ashby_rowe" / "ashby_rowe_budget_fy27.xlsx")


if __name__ == "__main__":
    nortech()
    ashby_rowe()
    ashby_rowe_budget()
    print(sorted(str(p.relative_to(OUT)) for p in OUT.rglob("*.*")))
