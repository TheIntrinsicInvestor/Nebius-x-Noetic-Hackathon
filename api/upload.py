"""User uploads: management accounts spreadsheet -> quarters, saved with the extracted covenants as a case."""
import calendar
import io
import json
import re
from datetime import date, datetime

from openpyxl import load_workbook

from api.extract_financials import MappedLine
from api.models import CovenantSet, Quarter

FLOWS = ("revenue", "cogs", "opex", "da", "interest", "capex", "principal", "addbacks")
BALANCES = ("cash", "debt", "receivables", "inventory", "payables")
UNITS = {"units": 1_000_000, "thousands": 1_000, "millions": 1}


def _month_end(v) -> date:
    if isinstance(v, (datetime, date)):
        y, m = v.year, v.month
    else:
        s = str(v).strip()
        for fmt in ("%b-%y", "%b %y", "%b-%Y", "%b %Y", "%Y-%m-%d", "%Y-%m", "%d/%m/%Y"):
            try:
                d = datetime.strptime(s, fmt)
                y, m = d.year, d.month
                break
            except ValueError:
                continue
        else:
            raise ValueError(f"column header '{s}' is not a month (use e.g. Sep-24 or 2024-09-30)")
    return date(y, m, calendar.monthrange(y, m)[1])


def read_sheet(data: bytes) -> tuple[list[date], list[tuple[str, list[float]]]]:
    """First column = line names, header row = month or quarter ends."""
    ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise ValueError("the spreadsheet needs a header row of dates and at least one account line")
    cols = [_month_end(v) for v in rows[0][1:] if v not in (None, "")]
    lines = []
    for r in rows[1:]:
        if not r or r[0] in (None, ""):
            continue
        vals = [float(v or 0) for v in r[1:1 + len(cols)]]
        lines.append((str(r[0]).strip(), vals))
    return cols, lines


def build_quarters(cols: list[date], lines: list[tuple[str, list[float]]], mapping: list[MappedLine],
                   unit: str = "thousands") -> list[Quarter]:
    """Sum flow lines over each calendar quarter; take balance lines at the quarter end. Engine works in millions."""
    if unit not in UNITS:
        raise ValueError(f"unit must be one of {list(UNITS)}")
    div = UNITS[unit]
    cat = {m.name: m.category for m in mapping}
    gap = (cols[1].year - cols[0].year) * 12 + cols[1].month - cols[0].month if len(cols) > 1 else 3
    per_q = 1 if gap >= 3 else 3
    groups: dict[date, list[int]] = {}
    for i, c in enumerate(cols):
        qm = ((c.month - 1) // 3 + 1) * 3
        groups.setdefault(date(c.year, qm, calendar.monthrange(c.year, qm)[1]), []).append(i)
    qs = []
    for qend, idx in sorted(groups.items()):
        if len(idx) < per_q:
            continue  # incomplete quarter
        tot = {k: 0.0 for k in FLOWS + BALANCES}
        for name, vals in lines:
            k = cat.get(name, "other")
            if k in FLOWS:
                tot[k] += sum(vals[i] for i in idx)
            elif k in BALANCES:
                tot[k] += vals[idx[-1]]
        t = {k: v / div for k, v in tot.items()}
        qs.append(Quarter(
            period_end=qend, revenue=t["revenue"], cogs=t["cogs"],
            ebit=t["revenue"] - t["cogs"] - t["opex"] - t["da"] - t["addbacks"], da=t["da"],
            interest_expense=t["interest"], capex=t["capex"], principal_paid=t["principal"], addbacks=t["addbacks"],
            cash=t["cash"], debt=t["debt"], receivables=t["receivables"], inventory=t["inventory"],
            payables=t["payables"]))
    if len(qs) < 4:
        raise ValueError(f"need at least 4 complete quarters of accounts to test covenants; found {len(qs)}")
    return qs


def slug(name: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return "u-" + (s or "company")[:40]


def save_case(root, case: str, qs: list[Quarter], cs: CovenantSet, meta: dict, mapping: list[MappedLine]) -> None:
    d = root / case
    d.mkdir(parents=True, exist_ok=True)
    (d / "quarters.json").write_text(json.dumps([q.model_dump(mode="json") for q in qs], indent=1), encoding="utf-8")
    (d / "covenants.json").write_text(cs.model_dump_json(indent=1), encoding="utf-8")
    (d / "meta.json").write_text(json.dumps(meta, indent=1), encoding="utf-8")
    (d / "mapping.json").write_text(json.dumps([m.model_dump() for m in mapping], indent=1), encoding="utf-8")
