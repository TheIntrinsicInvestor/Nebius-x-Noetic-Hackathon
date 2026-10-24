"""Messy management-account line names -> standard categories, via an open model."""
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook
from pydantic import BaseModel

from api.llm import PRIMARY_MODEL, CallStats, chat_json

Category = Literal["revenue", "cogs", "opex", "da", "interest", "capex", "cash", "debt", "receivables",
                   "inventory", "payables", "principal", "addbacks", "other"]


class MappedLine(BaseModel):
    name: str
    category: Category
    confidence: float


class _Mapped(BaseModel):
    lines: list[MappedLine]


SYSTEM = """You map a small company's management-account line names to standard categories used to test
bank covenants. Categories: revenue, cogs (direct costs), opex (overheads), da (depreciation/amortisation),
interest, capex, cash, debt (bank loans, RCF, overdraft), receivables (debtors), inventory (stock/WIP),
payables (creditors), principal (scheduled loan repayments), addbacks (exceptional/one-off items), other. Bookkeepers abbreviate (e.g. "Dep'n",
"int pd", "a/c"). Return every input line once, in order, with a confidence between 0 and 1."""


def read_xlsx_lines(path: Path) -> list[str]:
    ws = load_workbook(path, read_only=True).active
    return [str(r[0]) for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]]


def map_lines(names: list[str], model: str = PRIMARY_MODEL) -> tuple[list[MappedLine], CallStats]:
    data, stats = chat_json(model, SYSTEM, "\n".join(names), _Mapped.model_json_schema(), "mapped_lines")
    return _Mapped.model_validate(data).lines, stats
